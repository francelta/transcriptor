import os
import io
import datetime
import math
from rest_framework import status
from rest_framework.views import APIView
from rest_framework.response import Response
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .models import TranscriptionJob
from .serializers import TranscriptionJobSerializer
from .services.llm_processor import (
    process_transcript_with_llm,
    map_speakers_by_language_and_role,
    normalize_speaker_label
)

import gc
import threading
import json
import time
import logging
import requests
from django.conf import settings
from django.http import StreamingHttpResponse

logger = logging.getLogger(__name__)

# Constantes de ejecución optimizadas para CPU y estabilidad de memoria (16 GB RAM)
WHISPERX_DEVICE = getattr(settings, 'WHISPERX_DEVICE', 'cpu')
WHISPERX_BATCH_SIZE = getattr(settings, 'WHISPERX_BATCH_SIZE', 4)
# compute_type='int8' ejecuta el mismo modelo large-v3 con vectorización AVX2 de alta velocidad en CPU
WHISPERX_COMPUTE_TYPE = getattr(settings, 'WHISPERX_COMPUTE_TYPE', 'int8')
WHISPERX_MODEL_NAME = getattr(settings, 'WHISPERX_MODEL_NAME', 'large-v3')

# Control de cancelación en memoria
_active_jobs_cancel_flags = {}
_jobs_lock = threading.Lock()


def normalize_speaker_label(raw_speaker):
    """
    Normaliza el identificador de orador devuelto por la diarización
    a la nomenclatura estándar de la transcripción (I, R, R1, R2, etc.)
    o preserva etiquetas si ya vienen formateadas.
    """
    if not raw_speaker:
        return 'I'
    clean = str(raw_speaker).strip()
    if clean.startswith('SPEAKER_'):
        # Mapea SPEAKER_00 -> I, SPEAKER_01 -> R, SPEAKER_02 -> R2, etc.
        try:
            spk_num = int(clean.split('_')[-1])
            if spk_num == 0:
                return 'I'
            elif spk_num == 1:
                return 'R'
            else:
                return f'R{spk_num}'
        except Exception:
            return clean
    return clean


import threading
import json
import time
from django.http import StreamingHttpResponse

def append_job_log(job_id, message, level='info'):
    """Registra un evento con timestamp para el trabajo en la BD y en consola"""
    try:
        now_str = datetime.datetime.now().strftime("%H:%M:%S")
        log_entry = {
            'time': now_str,
            'message': message,
            'level': level
        }
        print(f"[{now_str}] [{level.upper()}] [Job {job_id}]: {message}")
        job = TranscriptionJob.objects.get(pk=job_id)
        current_logs = list(job.logs or [])
        current_logs.append(log_entry)
        job.logs = current_logs
        job.save(update_fields=['logs'])
    except Exception as e:
        print(f"Error guardando log para job {job_id}: {e}")


def stop_runpod_instance(job_id=None):
    """
    Ejecuta el apagado automático inmediato de la instancia de RunPod
    para congelar la facturación de GPU.
    """
    pod_id = getattr(settings, 'RUNPOD_POD_ID', '')
    api_key = getattr(settings, 'RUNPOD_API_KEY', '')

    if pod_id and api_key and api_key != 'PEGA_AQUÍ_TU_API_KEY':
        try:
            stop_url = f"https://rest.runpod.io/v1/pods/{pod_id}/stop"
            resp = requests.post(
                stop_url,
                headers={"Authorization": f"Bearer {api_key}"},
                timeout=5
            )
            msg = f"RunPod [{pod_id}] apagado automáticamente con éxito (Status: {resp.status_code}). Facturación congelada."
            if job_id:
                append_job_log(job_id, msg, "success")
            else:
                print(msg)
        except Exception as e:
            err = f"Error apagando RunPod [{pod_id}]: {e}"
            logger.warning(err)
            if job_id:
                append_job_log(job_id, err, "warning")
            else:
                print(err)
    else:
        msg = "Auto-apagado RunPod omitido: RUNPOD_POD_ID o RUNPOD_API_KEY no están configurados con una clave válida."
        if job_id:
            append_job_log(job_id, msg, "info")
        else:
            print(msg)


def execute_whisperx_pipeline(job_id, hf_token, min_speakers, max_speakers, llm_config=None):
    """
    Ejecuta la transcripción delegando de forma estricta y exclusiva en el worker remoto GPU de RunPod.
    Política estricta ZERO CPU FALLBACK: Prohibición absoluta de ejecutar modelos locales en CPU.
    Si el worker remoto o el túnel fallan, se aborta inmediatamente con error descriptivo.
    Timeout global de llamadas al worker: 900 segundos.
    """
    from .models import RunPodConfig
    from .services.runpod_orchestrator import RunPodOrchestrator

    try:
        job = TranscriptionJob.objects.get(pk=job_id)
        audio_path = job.audio_file.path
        config = RunPodConfig.get_solo()

        local_port = config.local_proxy_port or 8005
        pod_host = (config.pod_host or '').strip()
        pod_ssh_port = config.pod_ssh_port or 22
        ssh_key_path = config.ssh_key_path or '~/.ssh/id_rsa'
        effective_hf_token = (hf_token or config.hf_token or '').strip()

        orchestrator = RunPodOrchestrator()

        append_job_log(job_id, f"Verificando disponibilidad de worker remoto GPU (localhost:{local_port})...", "info")

        # Asegurar túnel y worker activo si hay configuración de host
        if pod_host:
            status_check = orchestrator.ensure_worker_ready(
                host=pod_host,
                ssh_port=pod_ssh_port,
                ssh_key_path=ssh_key_path,
                local_port=local_port,
                hf_token=effective_hf_token
            )
            if status_check.get("status") not in ["connected"]:
                err_msg = status_check.get("error", "No se pudo conectar o aprovisionar el worker GPU.")
                append_job_log(job_id, f"ERROR: Fallo conectando con el worker GPU: {err_msg}", "error")
                raise RuntimeError(f"Fallo de conexión con worker remoto GPU: {err_msg}. Transcripción abortada (Zero CPU Fallback).")
        else:
            # Si no hay pod_host configurado, verificar si al menos el proxy ya responde
            if not orchestrator.check_health(local_port):
                raise RuntimeError(
                    f"No hay host de Pod configurado y el worker en localhost:{local_port} no responde. "
                    "Por favor, configura la IP del Pod o detecta automáticamente antes de transcribir."
                )

        gpu_worker_url = f"http://127.0.0.1:{local_port}/transcribe"
        append_job_log(job_id, f"Worker GPU listo. Enviando audio a {gpu_worker_url} (Timeout: 900s)...", "info")

        content_type = 'audio/mpeg'
        ext = os.path.splitext(job.original_filename)[1].lower()
        if ext in ['.wav']:
            content_type = 'audio/wav'
        elif ext in ['.ogg']:
            content_type = 'audio/ogg'
        elif ext in ['.m4a']:
            content_type = 'audio/mp4'
        elif ext in ['.flac']:
            content_type = 'audio/flac'

        with open(audio_path, 'rb') as af:
            files = {'file': (job.original_filename, af.read(), content_type)}
        data = {
            'min_speakers': min_speakers,
            'max_speakers': max_speakers,
            'hf_token': effective_hf_token
        }

        append_job_log(job_id, f"Transcribiendo y diarizando en GPU RunPod (oradores: {min_speakers}-{max_speakers})...", "info")
        
        try:
            resp = requests.post(gpu_worker_url, files=files, data=data, timeout=900)
            if resp.status_code != 200:
                raise RuntimeError(f"El worker GPU remoto devolvió código de error HTTP {resp.status_code}: {resp.text}")
            worker_response = resp.json()
        except requests.Timeout:
            raise RuntimeError("La llamada de transcripción remota expiró tras 900 segundos (Timeout).")
        except requests.RequestException as req_err:
            raise RuntimeError(f"Error de comunicación HTTP con el worker GPU remoto ({gpu_worker_url}): {req_err}")

        append_job_log(job_id, "Transcripción GPU recibida exitosamente del worker de RunPod.", "success")

        # Auto-apagado opcional si está configurada la API de RunPod
        stop_runpod_instance(job_id=job_id)

        # Extraer segmentos y metadatos devueltos por el worker
        if isinstance(worker_response, dict):
            result = worker_response
            if 'duration' in worker_response:
                job.duration_seconds = round(float(worker_response['duration']), 2)
                job.save(update_fields=['duration_seconds'])
        elif isinstance(worker_response, list):
            result = {'segments': worker_response}
        else:
            result = {'segments': []}


        # -------------------------------------------------------------
        # CONSTRUCCIÓN DE BLOQUES COMPATIBLES CON KARAOKE Y FRONTEND
        # -------------------------------------------------------------
        blocks = []
        all_words = []
        block_idx = 1
        last_speaker = 'SPEAKER_00'

        raw_segments = result.get("segments", []) if isinstance(result, dict) else []
        for seg in raw_segments:
            text = seg.get("text", "").strip()
            if not text:
                continue

            raw_spk = seg.get("speaker") or last_speaker
            last_speaker = raw_spk

            seg_words = []
            for w in seg.get("words", []):
                w_text = w.get("word", "").strip()
                w_start = w.get("start")
                w_end = w.get("end")

                start_val = round(float(w_start), 2) if w_start is not None else round(float(seg.get("start", 0.0)), 2)
                end_val = round(float(w_end), 2) if w_end is not None else round(float(seg.get("end", start_val + 0.3)), 2)
                w_speaker = w.get("speaker", raw_spk)

                word_item = {
                    'word': w_text,
                    'start': start_val,
                    'end': end_val,
                    'speaker': w_speaker
                }
                seg_words.append(word_item)
                all_words.append(word_item)

            language = seg.get("language")
            blocks.append({
                'id': block_idx,
                'speaker': raw_spk,
                'language': language,
                'text': text,
                'start': round(float(seg.get("start", 0.0)), 2),
                'end': round(float(seg.get("end", 0.0)), 2),
                'words': seg_words
            })
            block_idx += 1

        # Mapeo inicial de oradores según rol e idioma (I: preguntas/inglés, R1/R2: español)
        initial_spk_map = map_speakers_by_language_and_role(blocks)
        for b in blocks:
            old_s = b.get('speaker')
            if old_s in initial_spk_map:
                b['speaker'] = initial_spk_map[old_s]
                for w in b.get('words', []):
                    w['speaker'] = initial_spk_map[old_s]

        # Si no se detectó duración previamente, inferir del último segmento
        if job.duration_seconds <= 0.0 and blocks:
            last_block = blocks[-1]
            job.duration_seconds = round(float(last_block.get('end', 0.0)), 2)
            job.save(update_fields=['duration_seconds'])

        # Comprobar cancelación antes de LLM
        with _jobs_lock:
            if _active_jobs_cancel_flags.get(str(job_id), False):
                append_job_log(job_id, "Proceso detenido a petición del usuario.", "warning")
                return

        # -------------------------------------------------------------
        # FASE 4: Capa semántica de Post-Procesado con LLM
        # -------------------------------------------------------------
        append_job_log(job_id, "Iniciando post-procesado semántico con LLM (roles de oradores, Spanglish, interrupciones)...", "info")
        processed_blocks = process_transcript_with_llm(blocks, hf_token=hf_token, llm_config=llm_config)

        # Sincronizar etiquetas de orador en words_data a partir de los bloques post-procesados
        for b in processed_blocks:
            spk = b.get('speaker', 'I')
            if 'words' in b:
                for w in b['words']:
                    w['speaker'] = spk

        job.transcript_blocks = processed_blocks
        job.words_data = all_words
        job.status = 'completed'
        job.save()
        append_job_log(job_id, "¡Transcripción y post-procesado semántico completados con éxito!", "success")

    except Exception as e:
        import traceback
        err_msg = str(e)
        trace_str = traceback.format_exc()
        print(f"Error crítico en transcripción {job_id}:\n{trace_str}")
        try:
            job = TranscriptionJob.objects.get(pk=job_id)
            job.status = 'failed'
            job.error_message = err_msg
            job.save(update_fields=['status', 'error_message'])
            append_job_log(job_id, f"ERROR: {err_msg}", "error")
        except Exception:
            pass



class ActiveJobsView(APIView):
    """
    GET /api/transcribe/active/
    Devuelve los trabajos actualmente en estado 'processing'.
    """
    def get(self, request):
        active_jobs = TranscriptionJob.objects.filter(status='processing')
        serializer = TranscriptionJobSerializer(active_jobs, many=True, context={'request': request})
        return Response({'count': active_jobs.count(), 'jobs': serializer.data})


class KillAllActiveJobsView(APIView):
    """
    POST /api/transcribe/kill-active/
    Cancela y mata todos los procesos en ejecución para dejar el servidor limpio.
    """
    def post(self, request):
        with _jobs_lock:
            active_jobs = TranscriptionJob.objects.filter(status='processing')
            killed_ids = []
            for j in active_jobs:
                _active_jobs_cancel_flags[str(j.id)] = True
                j.status = 'failed'
                j.error_message = 'Cancelado/detenido forzosamente por el usuario desde el panel.'
                j.save(update_fields=['status', 'error_message'])
                append_job_log(j.id, "Proceso detenido y liberado forzosamente desde el panel.", "warning")
                killed_ids.append(str(j.id))
        
        # Forzar recolección de basura
        gc.collect()
        return Response({
            'status': 'success',
            'killed_count': len(killed_ids),
            'killed_ids': killed_ids
        }, status=status.HTTP_200_OK)


class CancelJobView(APIView):
    """
    POST /api/transcribe/<id>/cancel/
    Cancela un trabajo en curso y limpia sus recursos.
    """
    def post(self, request, pk):
        job = get_object_or_404(TranscriptionJob, pk=pk)
        with _jobs_lock:
            _active_jobs_cancel_flags[str(job.id)] = True
        job.status = 'failed'
        job.error_message = 'Cancelado por el usuario.'
        job.save(update_fields=['status', 'error_message'])
        append_job_log(job.id, "Proceso cancelado por el usuario.", "warning")
        gc.collect()
        return Response({'status': 'cancelled', 'job_id': str(job.id)}, status=status.HTTP_200_OK)


class TranscribeAudioView(APIView):
    """
    POST /api/transcribe/
    Garantiza una ÚNICA ejecución activa a la vez: cancela cualquier proceso anterior
    antes de iniciar el nuevo trabajo.
    """
    def post(self, request):
        audio = request.FILES.get('audio_file')
        if not audio:
            return Response({'error': 'No audio_file provided'}, status=status.HTTP_400_BAD_REQUEST)

        task_number = request.data.get('task_number', 'Tilaus numero / Task ID')
        original_filename = audio.name
        current_date_str = datetime.datetime.now().strftime("%B %d, %Y")

        hf_token = (request.data.get('hf_token') or getattr(settings, 'HF_TOKEN', '')).strip()
        min_speakers = int(request.data.get('min_speakers', 2))
        max_speakers = int(request.data.get('max_speakers', 3))

        # Configuración opcional del LLM enviada desde el cliente
        llm_enabled = request.data.get('llm_enabled', True)
        if isinstance(llm_enabled, str):
            llm_enabled = llm_enabled.lower() in ['true', '1', 'yes']
        llm_api_url = request.data.get('llm_api_url') or getattr(settings, 'LLM_API_URL', '')
        llm_api_key = request.data.get('llm_api_key') or getattr(settings, 'LLM_API_KEY', '')
        llm_model = request.data.get('llm_model') or getattr(settings, 'LLM_MODEL', '')

        llm_config = {
            'enabled': llm_enabled,
            'api_url': llm_api_url,
            'api_key': llm_api_key,
            'model': llm_model
        }

        # -------------------------------------------------------------
        # REGLA ESTRICTA: Solo un trabajo a la vez
        # Cancelar y limpiar cualquier trabajo previo que estuviera en 'processing'
        # -------------------------------------------------------------
        with _jobs_lock:
            active_jobs = TranscriptionJob.objects.filter(status='processing')
            for old_job in active_jobs:
                _active_jobs_cancel_flags[str(old_job.id)] = True
                old_job.status = 'failed'
                old_job.error_message = 'Cancelado automáticamente al iniciar una nueva transcripción.'
                old_job.save(update_fields=['status', 'error_message'])
                append_job_log(old_job.id, "Trabajo abortado: se ha iniciado una nueva transcripción.", "warning")

        initial_log = {
            'time': datetime.datetime.now().strftime("%H:%M:%S"),
            'message': f"Archivo '{original_filename}' recibido. Iniciando delegación en Worker GPU RunPod (Zero CPU Fallback)...",
            'level': 'info'
        }


        job = TranscriptionJob.objects.create(
            audio_file=audio,
            original_filename=original_filename,
            status='processing',
            metadata_header={
                'task_number': task_number,
                'transcribed_date': current_date_str,
                'comments': '--'
            },
            transcript_blocks=[],
            words_data=[],
            logs=[initial_log],
            error_message=''
        )

        with _jobs_lock:
            _active_jobs_cancel_flags[str(job.id)] = False

        # Iniciar pipeline de WhisperX en hilo exclusivo en segundo plano con soporte LLM
        thread = threading.Thread(
            target=execute_whisperx_pipeline,
            args=(job.id, hf_token, min_speakers, max_speakers, llm_config),
            daemon=True
        )
        thread.start()

        serializer = TranscriptionJobSerializer(job, context={'request': request})
        return Response(serializer.data, status=status.HTTP_201_CREATED)


from django.views import View

class JobLogsStreamView(View):
    """
    GET /api/transcribe/<id>/events/
    Server-Sent Events (SSE) para emitir logs y cambios de estado en tiempo real.
    Hereda de View estándar para permitir 'Accept: text/event-stream' sin bloqueo de DRF.
    """
    def get(self, request, pk):
        job = get_object_or_404(TranscriptionJob, pk=pk)

        def event_stream():
            last_index = 0
            while True:
                try:
                    current_job = TranscriptionJob.objects.get(pk=pk)
                except TranscriptionJob.DoesNotExist:
                    break

                all_logs = current_job.logs or []
                if len(all_logs) > last_index:
                    for entry in all_logs[last_index:]:
                        payload = json.dumps({
                            'type': 'log',
                            'log': entry,
                            'status': current_job.status
                        })
                        yield f"data: {payload}\n\n"
                    last_index = len(all_logs)

                if current_job.status in ['completed', 'failed']:
                    # Emitir evento final con el objeto de trabajo actualizado
                    final_data = TranscriptionJobSerializer(current_job, context={'request': request}).data
                    final_payload = json.dumps({
                        'type': 'status_change',
                        'status': current_job.status,
                        'job': final_data
                    })
                    yield f"data: {final_payload}\n\n"
                    break

                time.sleep(0.5)

        response = StreamingHttpResponse(event_stream(), content_type='text/event-stream')
        response['Cache-Control'] = 'no-cache'
        response['X-Accel-Buffering'] = 'no'
        return response


class JobDetailView(APIView):
    """
    GET /api/transcribe/<id>/
    """
    def get(self, request, pk):
        job = get_object_or_404(TranscriptionJob, pk=pk)
        serializer = TranscriptionJobSerializer(job, context={'request': request})
        return Response(serializer.data)


class SaveTranscriptView(APIView):
    """
    PATCH /api/transcribe/<id>/save/
    Updates transcript_blocks and metadata_header.
    """
    def patch(self, request, pk):
        job = get_object_or_404(TranscriptionJob, pk=pk)
        transcript_blocks = request.data.get('transcript_blocks')
        metadata_header = request.data.get('metadata_header')

        if transcript_blocks is not None:
            job.transcript_blocks = transcript_blocks
        if metadata_header is not None:
            job.metadata_header = metadata_header

        job.save()
        serializer = TranscriptionJobSerializer(job, context={'request': request})
        return Response(serializer.data, status=status.HTTP_200_OK)


import re
import urllib.parse

class ExportDocxView(APIView):
    """
    GET /api/transcribe/<id>/export-docx/
    Genera un archivo .docx según las directrices exactas del cliente:
    - Nombre del archivo idéntico al original sustituyendo la extensión por .docx
    - Tipografía obligatoria: Verdana 8 pt en todo el documento
    - Cabecera estructurada idéntica
    - Línea en blanco obligatoria entre intervenciones de oradores
    - Renderizado de palabras con énfasis/subrayado (<u>...</u>) con run.font.underline = True
    """
    def get(self, request, pk):
        job = get_object_or_404(TranscriptionJob, pk=pk)
        
        doc = Document()

        # Configurar márgenes estándar (1 pulgada)
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # Configurar estilo Normal por defecto: Verdana 8 pt
        normal_style = doc.styles['Normal']
        normal_style.font.name = 'Verdana'
        normal_style.font.size = Pt(8)
        normal_style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Helper para crear párrafos estrictamente en Verdana 8 pt
        def add_p(text='', bold=False, underline=False, space_after=0, space_before=0):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(space_after)
            p.paragraph_format.space_before = Pt(space_before)
            p.paragraph_format.line_spacing = 1.0
            if text:
                r = p.add_run(text)
                r.font.name = 'Verdana'
                r.font.size = Pt(8)
                r.bold = bold
                r.underline = underline
            return p

        # Metadatos de la cabecera
        meta = job.metadata_header or {}
        task_no = meta.get('task_number', 'Tilaus numero / Task ID')
        filename = job.original_filename or 'recording.mp3'
        
        # Duración en minutos (entero o redondeo hacia arriba)
        duration_mins = math.ceil(job.duration_seconds / 60.0) if job.duration_seconds > 0 else 0
        length_str = f"{duration_mins} min" if duration_mins > 0 else "0 min"
        
        # Formato de fecha estilo inglés: "September 12, 2026"
        created_dt = job.created_at if job.created_at else datetime.datetime.now()
        formatted_date = created_dt.strftime("%B %d, %Y")
        comments = meta.get('comments', '--')

        # 1. Cabecera exacta obligatoria
        add_p(f"Number of transcription:        {task_no}")
        add_p(f"Name of recording file:         {filename}")
        add_p(f"Length of recording:            {length_str}")
        add_p(f"Transcribed on:                 {formatted_date}", space_after=12)

        add_p("Indicators used:     I: Interviewer(s)")
        add_p("                     R: Respondent(s)", space_after=12)

        add_p("Other notations:     ...         Interrupted or continued statement")
        add_p("                     (-)         Omitted word or part of word")
        add_p("                     (--)        Omitted part of speech")
        add_p("                     (word)      Unclear word or uncertain spelling")
        add_p("                     underlined  Word or part of speech with particular emphasis")
        add_p("                     [brackets]  Transcriber's comment")
        add_p("                     [os]        Overlapping speech")
        add_p("                     [qs]        Quiet speech")
        add_p("                     [us]        Unclear speech", space_after=12)

        add_p(f"Comments:            {comments}", space_after=18)

        # Salto / Separación antes de las intervenciones
        add_p("", space_after=0)

        # 2. Intervenciones de oradores con regla estricta:
        # UNA LÍNEA EN BLANCO ENTRE CADA INTERVENCIÓN CUANDO CAMBIA EL ORADOR
        blocks = job.transcript_blocks or []
        for i, block in enumerate(blocks):
            speaker = block.get('speaker', 'I').strip()
            text = block.get('text', '').strip()

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0

            # Etiqueta del orador (ej: "I: ", "R: ", "R1: ", "R2: ") en negrita Verdana 8pt
            label_run = p.add_run(f"{speaker}: ")
            label_run.font.name = 'Verdana'
            label_run.font.size = Pt(8)
            label_run.bold = True

            # Parsear texto respetando posibles etiquetas <u>...</u> para palabras con énfasis
            if '<u>' in text and '</u>' in text:
                parts = re.split(r'(<u>.*?</u>)', text, flags=re.IGNORECASE)
                for part in parts:
                    if not part:
                        continue
                    if part.lower().startswith('<u>') and part.lower().endswith('</u>'):
                        underlined_content = part[3:-4]
                        r = p.add_run(underlined_content)
                        r.font.name = 'Verdana'
                        r.font.size = Pt(8)
                        r.underline = True
                    else:
                        r = p.add_run(part)
                        r.font.name = 'Verdana'
                        r.font.size = Pt(8)
            else:
                r = p.add_run(text)
                r.font.name = 'Verdana'
                r.font.size = Pt(8)

            # Dejar obligatoriamente UNA LÍNEA EN BLANCO solo cuando el orador cambia respecto a la siguiente intervención
            if i < len(blocks) - 1:
                next_speaker = blocks[i + 1].get('speaker', 'I').strip()
                if next_speaker != speaker:
                    blank_p = doc.add_paragraph()
                    blank_p.paragraph_format.space_before = Pt(0)
                    blank_p.paragraph_format.space_after = Pt(0)
                    blank_p.paragraph_format.line_spacing = 1.0
                    blank_r = blank_p.add_run("")
                    blank_r.font.name = 'Verdana'
                    blank_r.font.size = Pt(8)

        # Generar nombre del archivo idéntico al audio subido sustituyendo la extensión por .docx
        base_name, _ = os.path.splitext(job.original_filename or 'transcription.mp3')
        export_filename = f"{base_name}.docx"

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        quoted_filename = urllib.parse.quote(export_filename)
        response['Content-Disposition'] = f'attachment; filename="{export_filename}"; filename*=UTF-8\'\'{quoted_filename}'
        return response


# =====================================================================
# Vistas de Configuración y Ciclo de Vida del Worker GPU Remoto / RunPod
# =====================================================================
# Vistas de Configuración y Ciclo de Vida del Worker GPU Remoto / RunPod
# =====================================================================

from .models import RunPodConfig
from .serializers import RunPodConfigSerializer
from .services.runpod_orchestrator import RunPodOrchestrator, query_runpod_graphql


class PodStatusView(APIView):
    """
    GET /api/pod/status/
    Devuelve si el puerto local 8005 o el worker remoto responde 200 OK.
    """
    def get(self, request):
        config = RunPodConfig.get_solo()
        orchestrator = RunPodOrchestrator()
        local_port = config.local_proxy_port or 8005

        is_listening = orchestrator.is_port_listening(port=local_port)
        is_healthy = orchestrator.check_health(local_port=local_port)

        if is_healthy:
            return Response({
                "status": "connected",
                "healthy": True,
                "local_proxy_port": local_port,
                "pod_host": config.pod_host,
                "pod_ssh_port": config.pod_ssh_port,
                "message": f"Worker GPU respondiendo 200 OK en localhost:{local_port}"
            }, status=status.HTTP_200_OK)
        else:
            return Response({
                "status": "disconnected",
                "healthy": False,
                "local_proxy_port": local_port,
                "pod_host": config.pod_host,
                "pod_ssh_port": config.pod_ssh_port,
                "message": f"Worker GPU no responde en localhost:{local_port}"
            }, status=status.HTTP_200_OK)


class WorkerConfigView(APIView):
    """
    GET /api/worker/config/
    POST /api/worker/config/
    Permite consultar y actualizar los parámetros de configuración dinámica del Pod.
    """
    def get(self, request):
        config = RunPodConfig.get_solo()
        serializer = RunPodConfigSerializer(config)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def post(self, request):
        config = RunPodConfig.get_solo()
        serializer = RunPodConfigSerializer(config, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_200_OK)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class WorkerDetectView(APIView):
    """
    POST /api/worker/detect/
    Consulta la API GraphQL de RunPod para autocompletar la IP y puerto SSH a partir del POD_ID y RUNPOD_API_KEY.
    """
    def post(self, request):
        config = RunPodConfig.get_solo()
        pod_id = request.data.get('pod_id') or config.pod_id
        api_key = request.data.get('runpod_api_key') or config.runpod_api_key

        if not pod_id or not api_key:
            return Response(
                {"error": "Debes especificar tanto Pod ID como RunPod API Key para la detección automática."},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            pod_info = query_runpod_graphql(pod_id=pod_id, api_key=api_key)
            config.pod_id = pod_id
            config.runpod_api_key = api_key
            config.pod_host = pod_info["pod_host"]
            config.pod_ssh_port = pod_info["pod_ssh_port"]
            config.save()

            return Response({
                "status": "success",
                "pod_id": pod_id,
                "pod_name": pod_info.get("pod_name"),
                "pod_host": pod_info["pod_host"],
                "pod_ssh_port": pod_info["pod_ssh_port"],
                "desired_status": pod_info.get("desired_status"),
                "message": f"IP ({pod_info['pod_host']}) y Puerto SSH ({pod_info['pod_ssh_port']}) detectados exitosamente."
            }, status=status.HTTP_200_OK)
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)


class WorkerHealthView(APIView):
    """
    GET /api/worker/health/
    Verifica el estado del worker remoto:
    - 200 OK si localhost:{LOCAL_PROXY_PORT} responde OK
    """
    def get(self, request):
        config = RunPodConfig.get_solo()
        orchestrator = RunPodOrchestrator()
        local_port = config.local_proxy_port or 8005

        is_listening = orchestrator.is_port_listening(port=local_port)
        is_healthy = orchestrator.check_health(local_port=local_port)

        if is_healthy:
            return Response({
                "status": "connected",
                "tunnel_active": True,
                "local_proxy_port": local_port,
                "pod_host": config.pod_host,
                "pod_ssh_port": config.pod_ssh_port,
                "message": "Worker GPU conectado y respondiendo 200 OK."
            }, status=status.HTTP_200_OK)

        if config.pod_host and (request.query_params.get('auto_connect', 'false').lower() in ['1', 'true']):
            res = orchestrator.ensure_worker_ready(
                host=config.pod_host,
                ssh_port=config.pod_ssh_port,
                ssh_key_path=config.ssh_key_path,
                local_port=local_port,
                hf_token=config.hf_token
            )
            return Response(res, status=status.HTTP_200_OK if res.get("status") == "connected" else status.HTTP_502_BAD_GATEWAY)

        return Response({
            "status": "disconnected",
            "tunnel_active": is_listening,
            "local_proxy_port": local_port,
            "pod_host": config.pod_host,
            "pod_ssh_port": config.pod_ssh_port,
            "message": "Worker GPU desconectado o no responde en el puerto proxy local."
        }, status=status.HTTP_200_OK)


class WorkerConnectView(APIView):
    """
    POST /api/worker/connect/ o /api/pod/connect/
    Ejecuta el flujo completo de conexión, túnel SSH y auto-aprovisionamiento del pod.
    """
    def post(self, request):
        config = RunPodConfig.get_solo()

        pod_host = (request.data.get('pod_host') or config.pod_host or '').strip()
        pod_ssh_port = request.data.get('pod_ssh_port') or config.pod_ssh_port or 22
        ssh_key_path = request.data.get('ssh_key_path') or config.ssh_key_path or '~/.ssh/id_rsa'
        hf_token = (request.data.get('hf_token') or config.hf_token or '').strip()
        local_port = int(request.data.get('local_proxy_port') or config.local_proxy_port or 8005)

        if not pod_host:
            return Response({"error": "No se ha especificado POD_HOST."}, status=status.HTTP_400_BAD_REQUEST)

        orchestrator = RunPodOrchestrator()
        result = orchestrator.ensure_worker_ready(
            host=pod_host,
            ssh_port=int(pod_ssh_port),
            ssh_key_path=ssh_key_path,
            local_port=local_port,
            hf_token=hf_token
        )

        http_status = status.HTTP_200_OK if result.get("status") == "connected" else status.HTTP_502_BAD_GATEWAY
        return Response(result, status=http_status)


class WorkerDisconnectView(APIView):
    """
    POST /api/worker/disconnect/
    Cierra el túnel SSH hacia el pod.
    """
    def post(self, request):
        orchestrator = RunPodOrchestrator()
        orchestrator.stop_tunnel()
        return Response({"status": "disconnected", "message": "Túnel SSH detenido."}, status=status.HTTP_200_OK)


# =====================================================================
# APROVISIONAMIENTO AUTOMÁTICO 1-CLIC: PROVISION / EVENTS / TERMINATE
# =====================================================================

# Estado global del aprovisionamiento (en memoria, accesible entre vistas)
_provision_lock = threading.Lock()
# Cola en memoria: canal de comunicación entre el hilo de aprovisionamiento y el SSE
# Evita el problema de visibilidad SQLite entre hilos en Django dev server
import queue as _queue_module
_provision_queue: _queue_module.Queue = _queue_module.Queue()
_provision_active: bool = False  # Flag para saber si hay un aprovisionamiento en curso

def _run_provision_flow(config_id: int, api_key: str, hf_token: str, ssh_key_path: str, local_port: int):
    """
    Ejecuta el flujo completo de aprovisionamiento en un thread de background.
    Comunica el progreso mediante _provision_queue (en memoria) para el SSE,
    y escribe el estado final en la DB.
    """
    global _provision_active
    from django.db import connection as _db_connection
    # Cerrar conexión heredada del proceso padre para evitar stale SQLite en el hilo
    _db_connection.close()

    from .models import RunPodConfig
    from .services.runpod_service import provision_pod, poll_pod_until_running
    from .services.runpod_orchestrator import RunPodOrchestrator, POD_SERVER_PY_CONTENT

    def push_log(step: str, message: str, level: str = "info"):
        """Emite un log al SSE (via queue) Y lo persiste en DB."""
        entry = {
            "step": step,
            "message": message,
            "level": level,
            "time": datetime.datetime.now().strftime("%H:%M:%S"),
        }
        # 1. Emitir a la queue (SSE la lee inmediatamente)
        _provision_queue.put(entry)
        # 2. Persistir en DB (para recuperación tras recarga)
        try:
            cfg = RunPodConfig.objects.get(pk=config_id)
            logs = list(cfg.provision_logs or [])
            logs.append(entry)
            cfg.provision_logs = logs
            cfg.save(update_fields=["provision_logs"])
        except Exception as ex:
            logger.warning(f"push_log DB write failed (non-critical): {ex}")

    def set_status(s: str):
        try:
            RunPodConfig.objects.filter(pk=config_id).update(provision_status=s)
        except Exception as ex:
            logger.warning(f"set_status DB write failed: {ex}")
        # Emitir marcador de estado a la queue para que el SSE lo detecte
        _provision_queue.put({"__status__": s})

    try:
        # ── PASO 1: Buscar todas las GPUs asequibles (lista priorizada) ──────
        set_status("provisioning")
        push_log("searching_gpu", "🔍 Buscando GPUs disponibles por debajo de 0.40 $/h en RunPod...", "info")

        from .services.runpod_service import find_all_affordable_gpus
        gpu_candidates = find_all_affordable_gpus(api_key)

        if not gpu_candidates:
            raise RuntimeError(
                "No hay ninguna GPU On-Demand disponible por debajo de 0.40 $/h en RunPod. "
                "Verifica la disponibilidad o ajusta el límite de precio."
            )

        gpu_names = ", ".join(f"{g['displayName']} ({g['price']:.3f}$/h)" for g in gpu_candidates)
        push_log(
            "searching_gpu",
            f"✅ {len(gpu_candidates)} GPU(s) candidata(s) encontrada(s): {gpu_names}",
            "success"
        )

        # ── PASO 2: Intentar crear pod con cada GPU por orden de prioridad ──
        gpu_info = None
        pod_id = None
        last_deploy_error = None

        for candidate in gpu_candidates:
            push_log(
                "creating_pod",
                f"🚀 Intentando {candidate['displayName']} @ {candidate['price']:.3f} $/h "
                f"({candidate['memoryInGb']} GB VRAM)...",
                "info"
            )
            try:
                pod_id = provision_pod(api_key, candidate["id"])
                gpu_info = candidate
                push_log("creating_pod", f"✅ Pod creado: {pod_id} con {gpu_info['displayName']}", "success")
                break  # Éxito — salir del bucle
            except RuntimeError as deploy_err:
                last_deploy_error = str(deploy_err)
                err_short = last_deploy_error[:120]
                push_log(
                    "creating_pod",
                    f"⚠️ {candidate['displayName']} sin instancias libres: {err_short}. Probando siguiente...",
                    "info"
                )
                continue

        if not pod_id or not gpu_info:
            raise RuntimeError(
                f"Ninguna GPU tuvo instancias disponibles en este momento. "
                f"Último error: {last_deploy_error or 'desconocido'}. "
                "Espera unos minutos y vuelve a intentarlo, o consulta la disponibilidad en RunPod."
            )

        # Guardar info de GPU y pod_id
        RunPodConfig.objects.filter(pk=config_id).update(
            provision_gpu_info=gpu_info,
            provisioned_pod_id=pod_id,
            pod_id=pod_id
        )

        # ── PASO 3: Esperar IP pública (polling cada 3s) ──────────────────────
        push_log("waiting_ip", f"⏳ Esperando que el pod {pod_id} alcance estado RUNNING (polling 3s)...", "info")
        pod_info = poll_pod_until_running(api_key, pod_id, max_retries=80, interval_seconds=3.0)

        pod_host = pod_info["pod_host"]
        pod_ssh_port = pod_info["pod_ssh_port"]
        push_log("waiting_ip", f"✅ IP pública obtenida: {pod_host}:{pod_ssh_port}", "success")

        # Persistir host y puerto SSH
        RunPodConfig.objects.filter(pk=config_id).update(
            pod_host=pod_host,
            pod_ssh_port=pod_ssh_port
        )

        # ── PASOS 4-7: Aprovisionamiento vía SSH (Paramiko) ──────────────────
        push_log("ssh_connect", f"🔗 Conectando por SSH a root@{pod_host}:{pod_ssh_port}...", "info")

        import paramiko, io as _io

        expanded_key = os.path.expanduser(ssh_key_path)
        ssh = paramiko.SSHClient()
        ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())

        # Reintentos de SSH (el pod puede tardar en arrancar el servicio)
        ssh_connected = False
        for ssh_attempt in range(12):
            try:
                ssh.connect(
                    hostname=pod_host,
                    port=int(pod_ssh_port),
                    username="root",
                    key_filename=expanded_key,
                    timeout=20,
                    banner_timeout=25,
                )
                ssh_connected = True
                if ssh.get_transport():
                    ssh.get_transport().set_keepalive(15)
                break
            except Exception as ssh_err:
                logger.warning(f"SSH intento {ssh_attempt+1}/12: {ssh_err}")
                time.sleep(5)

        if not ssh_connected:
            raise RuntimeError(
                f"No se pudo conectar por SSH a root@{pod_host}:{pod_ssh_port} tras 12 intentos. "
                "Verifica que la clave pública está registrada en RunPod Settings → SSH Public Keys."
            )

        push_log("ssh_connect", "✅ Conexión SSH establecida con éxito.", "success")

        def exec_cmd(cmd, timeout=300):
            stdin, stdout, stderr = ssh.exec_command(cmd, timeout=timeout)
            exit_status = stdout.channel.recv_exit_status()
            out = stdout.read().decode("utf-8", errors="replace").strip()
            err = stderr.read().decode("utf-8", errors="replace").strip()
            return exit_status, out, err

        # ── PASO 5: Instalar dependencias del sistema ─────────────────────────
        push_log("installing_deps", "📦 Verificando e instalando dependencias del sistema (ffmpeg, tmux, libsndfile1)...", "info")
        st, _, _ = exec_cmd("which tmux && which ffmpeg && which curl")
        if st != 0:
            exec_cmd(
                "apt-get update -qq && apt-get install -y -qq ffmpeg libsndfile1 tmux curl",
                timeout=180
            )

        # Verificar dependencias Python (WhisperX oficial + FastAPI)
        st, _, _ = exec_cmd("python3 -c 'import whisperx, soundfile, fastapi, uvicorn'")
        if st != 0:
            push_log("installing_deps", "📦 Instalando librerías Python (whisperx, pyannote.audio, fastapi, uvicorn, soundfile)...", "info")
            exec_cmd(
                "pip install --no-cache-dir whisperx fastapi uvicorn soundfile python-multipart scipy",
                timeout=420
            )
        push_log("installing_deps", "✅ Dependencias del sistema y Python verificadas.", "success")

        # ── PASO 6: Subir /server.py vía SFTP ────────────────────────────────
        push_log("uploading_script", "📤 Subiendo script de inferencia /server.py al pod...", "info")
        # Asegurar que el transporte SSH sigue vivo tras descargas largas
        if not ssh.get_transport() or not ssh.get_transport().is_active():
            ssh = paramiko.SSHClient()
            ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
            ssh.connect(
                hostname=pod_host,
                port=int(pod_ssh_port),
                username="root",
                key_filename=expanded_key,
                timeout=20,
            )
            if ssh.get_transport():
                ssh.get_transport().set_keepalive(15)

        sftp = ssh.open_sftp()
        with sftp.file("/server.py", "w") as remote_f:
            remote_f.write(POD_SERVER_PY_CONTENT)
        sftp.close()
        push_log("uploading_script", "✅ /server.py subido correctamente.", "success")

        # ── PASO 7: Iniciar worker en sesión tmux ─────────────────────────────
        push_log("starting_worker", "⚙️ Lanzando servidor GPU en sesión tmux 'worker'...", "info")
        exec_cmd("tmux kill-session -t worker 2>/dev/null || true")
        token_env = f"export HF_TOKEN='{hf_token.strip()}';" if hf_token else ""
        launch_cmd = (
            f"tmux new-session -d -s worker "
            f"'{token_env} python3 /server.py > /tmp/server.log 2>&1'"
        )
        exec_cmd(launch_cmd)
        ssh.close()
        push_log("starting_worker", "✅ Worker GPU lanzado en sesión tmux 'worker'.", "success")

        # ── PASO 8: Abrir túnel SSH ───────────────────────────────────────────
        push_log("opening_tunnel", f"🔒 Abriendo túnel SSH localhost:{local_port} → pod:8000...", "info")
        orchestrator = RunPodOrchestrator()
        orchestrator.start_tunnel(pod_host, pod_ssh_port, ssh_key_path, local_port=local_port, remote_port=8000)
        push_log("opening_tunnel", f"✅ Túnel SSH activo en localhost:{local_port}.", "success")

        # ── PASO 9: Healthcheck (polling hasta 200 OK) ────────────────────────
        push_log("healthcheck", f"🏥 Verificando que el worker responde en localhost:{local_port}/health...", "info")
        worker_ready = False
        for hc_attempt in range(30):  # Hasta 90 segundos
            time.sleep(3)
            try:
                r = requests.get(f"http://127.0.0.1:{local_port}/health", timeout=4)
                if r.status_code == 200:
                    worker_ready = True
                    break
            except Exception:
                pass

        if not worker_ready:
            raise RuntimeError(
                f"El worker GPU no respondió 200 OK en localhost:{local_port}/health tras 90 segundos. "
                "Revisa /tmp/server.log en el pod para más detalles."
            )

        # ── PASO 10: Listo ────────────────────────────────────────────────────
        set_status("ready")
        push_log(
            "ready",
            f"🎉 ¡GPU lista para transcribir! Worker responde en localhost:{local_port} "
            f"({gpu_info['displayName']} @ {gpu_info['price']:.3f} $/h).",
            "success"
        )

    except Exception as exc:
        logger.error(f"Error en flujo de aprovisionamiento: {exc}")
        import traceback
        traceback.print_exc()
        try:
            push_log("error", f"❌ Error: {exc}", "error")
            set_status("error")
        except Exception:
            pass
    finally:
        _provision_active = False


class PodProvisionView(APIView):
    """
    POST /api/pod/provision/
    Lanza el flujo completo de aprovisionamiento automático 1-clic en un thread de background:
    busca GPU económica → crea pod → espera IP → SSH → instala deps → sube server.py → tmux → túnel → healthcheck.
    Responde inmediatamente con 202 Accepted. El progreso se puede seguir via SSE en /api/pod/provision/events/
    """
    def post(self, request):
        global _provision_active
        from .models import RunPodConfig

        config = RunPodConfig.get_solo()

        # Verificar que hay API key configurada
        api_key = (request.data.get("runpod_api_key") or config.runpod_api_key or "").strip()
        if not api_key:
            return Response(
                {"error": "RunPod API Key no configurada. Ve a Configuración → Worker Config y añade tu clave."},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Si ya está aprovisionando ACTIVAMENTE (hilo vivo), evitar doble arranque.
        # Nota: _provision_active se resetea a False cuando el proceso Django reinicia,
        # evitando que un status "provisioning" antiguo en DB bloquee nuevas solicitudes.
        if _provision_active and config.provision_status == "provisioning":
            return Response(
                {"error": "Ya hay un aprovisionamiento en curso. Espera a que finalice o recarga la página."},
                status=status.HTTP_409_CONFLICT
            )

        # Si el status estaba atascado en "provisioning" pero no hay hilo activo
        # (reinicio del servidor), resetearlo para permitir nuevo intento
        if not _provision_active and config.provision_status == "provisioning":
            config.provision_status = "error"
            config.provision_logs = (config.provision_logs or []) + [{
                "step": "error",
                "message": "⚠️ Aprovisionamiento anterior interrumpido por reinicio del servidor. Puedes reintentar.",
                "level": "info",
                "time": datetime.datetime.now().strftime("%H:%M:%S"),
            }]
            config.save(update_fields=["provision_status", "provision_logs"])

        hf_token = (request.data.get("hf_token") or config.hf_token or "").strip()
        ssh_key_path = (request.data.get("ssh_key_path") or config.ssh_key_path or "~/.ssh/id_rsa").strip()
        local_port = int(request.data.get("local_proxy_port") or config.local_proxy_port or 8005)

        # Guardar la api_key si viene del request (sin campos manuales)
        if request.data.get("runpod_api_key"):
            config.runpod_api_key = api_key
            config.save(update_fields=["runpod_api_key"])

        # Limpiar queue antes de iniciar (descartar mensajes de runs anteriores)
        while not _provision_queue.empty():
            try:
                _provision_queue.get_nowait()
            except _queue_module.Empty:
                break

        config.provision_status = "provisioning"
        config.provision_logs = []
        config.provision_gpu_info = {}
        config.save(update_fields=["provision_status", "provision_logs", "provision_gpu_info"])

        _provision_active = True

        # Lanzar el flujo en background
        thread = threading.Thread(
            target=_run_provision_flow,
            args=(config.id, api_key, hf_token, ssh_key_path, local_port),
            daemon=True,
        )
        thread.start()

        return Response(
            {
                "status": "provisioning",
                "message": "Aprovisionamiento automático iniciado. Sigue el progreso en tiempo real via SSE.",
            },
            status=status.HTTP_202_ACCEPTED,
        )


class PodProvisionEventsView(View):
    """
    GET /api/pod/provision/events/
    Server-Sent Events (SSE). Lee desde _provision_queue (en memoria) para
    comunicación instantanea con el hilo de aprovisionamiento.
    Evita el problema de visibilidad SQLite entre hilos en el Django dev server.
    """
    def get(self, request):
        from django.db import connection as _db_conn

        def event_stream():
            # Leer logs históricos de DB (en caso de recarga / reconexion)
            try:
                _db_conn.close()  # Forzar nueva conexión limpia
                from .models import RunPodConfig
                cfg = RunPodConfig.objects.get(pk=1)
                historic_logs = list(cfg.provision_logs or [])
                current_status = cfg.provision_status
                gpu_info = cfg.provision_gpu_info or {}
            except Exception:
                historic_logs = []
                current_status = "provisioning"
                gpu_info = {}

            # Emitir logs históricos (para reconexiones)
            for entry in historic_logs:
                payload = json.dumps({
                    "type": "provision_log",
                    "step": entry.get("step", ""),
                    "message": entry.get("message", ""),
                    "level": entry.get("level", "info"),
                    "time": entry.get("time", ""),
                    "provision_status": current_status,
                    "gpu_info": gpu_info,
                })
                yield f"data: {payload}\n\n"

            if current_status in ("ready", "error"):
                # Ya terminó antes de que llegara el SSE: emitir evento final
                _db_conn.close()
                from .models import RunPodConfig
                try:
                    cfg = RunPodConfig.objects.get(pk=1)
                except Exception:
                    return
                final_payload = json.dumps({
                    "type": "provision_complete",
                    "provision_status": cfg.provision_status,
                    "pod_host": cfg.pod_host,
                    "pod_ssh_port": cfg.pod_ssh_port,
                    "provisioned_pod_id": cfg.provisioned_pod_id,
                    "gpu_info": cfg.provision_gpu_info or {},
                })
                yield f"data: {final_payload}\n\n"
                return

            # Escuchar la queue en tiempo real (sin poll a DB)
            max_wait_seconds = 600  # 10 minutos máx
            deadline = time.time() + max_wait_seconds

            while time.time() < deadline:
                try:
                    item = _provision_queue.get(timeout=2.0)
                except _queue_module.Empty:
                    # Heartbeat para mantener la conexión viva
                    yield ": heartbeat\n\n"
                    continue

                # Marcador de estado (sent by set_status)
                if "__status__" in item:
                    new_status = item["__status__"]
                    if new_status in ("ready", "error"):
                        # Emitir evento final con datos de DB
                        time.sleep(0.3)  # Dejar que la DB se actualice
                        _db_conn.close()
                        from .models import RunPodConfig
                        try:
                            cfg = RunPodConfig.objects.get(pk=1)
                        except Exception:
                            break
                        final_payload = json.dumps({
                            "type": "provision_complete",
                            "provision_status": new_status,
                            "pod_host": cfg.pod_host,
                            "pod_ssh_port": cfg.pod_ssh_port,
                            "provisioned_pod_id": cfg.provisioned_pod_id,
                            "gpu_info": cfg.provision_gpu_info or {},
                        })
                        yield f"data: {final_payload}\n\n"
                        break
                    continue

                # Entrada de log normal
                payload = json.dumps({
                    "type": "provision_log",
                    "step": item.get("step", ""),
                    "message": item.get("message", ""),
                    "level": item.get("level", "info"),
                    "time": item.get("time", ""),
                    "provision_status": "provisioning",
                    "gpu_info": {},
                })
                yield f"data: {payload}\n\n"

        response = StreamingHttpResponse(event_stream(), content_type="text/event-stream")
        response["Cache-Control"] = "no-cache"
        response["X-Accel-Buffering"] = "no"
        response["Connection"] = "keep-alive"
        return response


class PodTerminateView(APIView):
    """
    POST /api/pod/terminate/
    Elimina definitivamente el pod activo en RunPod (provisioned_pod_id o pod_id).
    Detiene el túnel SSH y resetea el estado de aprovisionamiento.
    IMPORTANTE: Esta acción CONGELA la facturación de GPU de forma inmediata.
    """
    def post(self, request):
        from .models import RunPodConfig
        from .services.runpod_service import terminate_pod
        from .services.runpod_orchestrator import RunPodOrchestrator

        config = RunPodConfig.get_solo()
        api_key = (config.runpod_api_key or "").strip()

        # Elegir el pod a terminar: primero el aprovisionado automáticamente, luego el manual
        pod_id = (
            request.data.get("pod_id")
            or config.provisioned_pod_id
            or config.pod_id
        ).strip()

        if not pod_id:
            return Response(
                {"error": "No hay ningún Pod ID configurado para terminar."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not api_key:
            return Response(
                {"error": "RunPod API Key no configurada. No se puede terminar el pod."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # Detener el túnel SSH primero
        try:
            orchestrator = RunPodOrchestrator()
            orchestrator.stop_tunnel()
        except Exception as e:
            logger.warning(f"Error deteniendo túnel SSH antes de terminar pod: {e}")

        # Terminar el pod en RunPod
        try:
            terminate_pod(api_key, pod_id)
        except Exception as exc:
            return Response(
                {"error": f"Error terminando el pod {pod_id} en RunPod: {exc}"},
                status=status.HTTP_502_BAD_GATEWAY,
            )

        # Resetear configuración
        global _provision_active
        _provision_active = False
        config.provisioned_pod_id = ""
        config.pod_host = ""
        config.pod_ssh_port = 22
        config.provision_status = "idle"
        config.provision_logs = []
        config.provision_gpu_info = {}
        config.save()

        return Response(
            {
                "status": "terminated",
                "pod_id": pod_id,
                "message": f"Pod {pod_id} eliminado correctamente. Facturación de GPU congelada.",
            },
            status=status.HTTP_200_OK,
        )
